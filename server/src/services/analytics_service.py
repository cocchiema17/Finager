import io
from datetime import date
from typing import Optional
from uuid import UUID
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.orm import Session
from fastapi import HTTPException, status
from src.repositories.space_repository import SpaceRepository
from src.repositories.analytics_repository import AnalyticsRepository


class AnalyticsService:
    @staticmethod
    def get_charts(
        db: Session,
        user_id: UUID,
        space_id: int,
        from_date: Optional[date] = None,
        to_date: Optional[date] = None,
    ) -> dict:
        # Verifica autorizzazione su Space (uguale al controllo if (!spaces.find(s => s.id == spaceId)))
        space = SpaceRepository.get_by_id(db, space_id)
        if not space or space.userId != user_id:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Space not found",
            )

        bar_chart = AnalyticsRepository.get_bar_chart_data(db, space_id, from_date, to_date)
        line_chart = AnalyticsRepository.get_line_chart_data(db, space_id, from_date, to_date)
        pie_chart = AnalyticsRepository.get_pie_chart_data(db, space_id, from_date, to_date)

        return {
            "barChart": bar_chart,
            "lineChart": line_chart,
            "pieChart": pie_chart,
        }

    @staticmethod
    def generate_excel_report(db: Session, user_id: UUID) -> io.BytesIO:
        data = AnalyticsRepository.get_report_data(db, user_id)

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        headers = ["Title", "Description", "Amount", "Category", "Date"]
        ws.append(headers)

        # Stile testata (sfondo scuro e testo bianco in grassetto)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Righe dati
        for row in data:
            ws.append([
                row["Title"],
                row["Description"],
                row["Amount"],
                row["Category"],
                row["Date"],
            ])

        # Adattamento larghezza colonne
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output