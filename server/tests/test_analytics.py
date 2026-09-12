import io
from openpyxl import load_workbook


def test_charts_and_report(client):
    # 1. Registra utente
    reg_res = client.post(
        "/api/auth/register",
        json={
            "firstName": "Alessandro",
            "lastName": "Manzoni",
            "email": "alessandro@example.com",
            "password": "password123",
        },
    )
    token = reg_res.cookies.get("session")
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Crea space
    s_res = client.post("/api/spaces", json={"name": "Budget Mensile"}, headers=headers)
    space_id = s_res.json()["id"]

    # 3. Inserisci due transazioni (una entrata e una uscita)
    client.post(
        "/api/transactions",
        json={
            "title": "Stipendio",
            "description": "Accredito mensile",
            "date": "2026-03-01",
            "categoryName": "Lavoro",
            "color": "#00FF00",
            "spaceId": space_id,
            "value": 1500.0,
        },
        headers=headers,
    )
    client.post(
        "/api/transactions",
        json={
            "title": "Bolletta Luce",
            "description": "Enel",
            "date": "2026-03-05",
            "categoryName": "Utenze",
            "color": "#FF0000",
            "spaceId": space_id,
            "value": -120.0,
        },
        headers=headers,
    )

    # 4. Test GET /api/charts
    charts_res = client.get(f"/api/charts?spaceId={space_id}", headers=headers)
    assert charts_res.status_code == 200, charts_res.text
    charts_data = charts_res.json()

    assert "barChart" in charts_data
    assert "lineChart" in charts_data
    assert "pieChart" in charts_data

    # Verifica barChart (2 tipi: revenue ed expense)
    assert len(charts_data["barChart"]) == 2

    # Verifica pieChart (2 categorie)
    assert len(charts_data["pieChart"]) == 2

    # Verifica lineChart (saldo cumulativo: 1500 - 120 = 1380.0)
    assert len(charts_data["lineChart"]) == 1
    assert charts_data["lineChart"][0]["value"] == 1380.0

    # 5. Test GET /api/charts con space non autorizzato / inesistente
    err_charts = client.get("/api/charts?spaceId=99999", headers=headers)
    assert err_charts.status_code == 400

    # 6. Test GET /api/report (Download Excel)
    report_res = client.get("/api/report", headers=headers)
    assert report_res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in report_res.headers["content-type"]

    # Verifica che il buffer sia un file Excel leggibile e contenga le righe corrette
    excel_file = io.BytesIO(report_res.content)
    wb = load_workbook(excel_file)
    ws = wb.active
    # Riga 1 è l'header, righe 2 e 3 sono le transazioni
    assert ws.max_row == 3
    assert ws.cell(row=1, column=1).value == "Title"